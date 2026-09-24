import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Style configuration — Global Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []

    def flush_code_block():
        nonlocal code_lines, in_code_block
        if not code_lines:
            in_code_block = False
            return
        code_text = "".join(code_lines).strip('\n')
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "F4F5F7")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
        code_lines = []
        in_code_block = False

    def flush_table():
        nonlocal table_rows, in_table
        if not table_rows:
            in_table = False
            return
        
        parsed_rows = []
        for r in table_rows:
            if re.match(r'^\s*\|?[\s\-:|]+\|?\s*$', r):
                continue
            cells = [c.strip() for c in r.strip().strip('|').split('|')]
            if cells:
                parsed_rows.append(cells)

        if not parsed_rows:
            table_rows = []
            in_table = False
            return

        num_cols = max(len(r) for r in parsed_rows)
        tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

        for row_idx, row_data in enumerate(parsed_rows):
            for col_idx in range(num_cols):
                cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                cell = tbl.cell(row_idx, col_idx)
                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)

                if row_idx == 0:
                    set_cell_background(cell, "2B4C7E")
                    run = p.add_run(re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text))
                    run.bold = True
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(10.5)
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                else:
                    if row_idx % 2 == 1:
                        set_cell_background(cell, "F9FAFB")
                    else:
                        set_cell_background(cell, "FFFFFF")
                    add_formatted_text(p, cell_text, default_size=10.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        table_rows = []
        in_table = False

    def add_formatted_text(paragraph, text, default_size=11):
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))', text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**'):
                r = paragraph.add_run(token[2:-2])
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(default_size)
            elif token.startswith('*') and token.endswith('*'):
                r = paragraph.add_run(token[1:-1])
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(default_size)
            elif token.startswith('`') and token.endswith('`'):
                r = paragraph.add_run(token[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(default_size - 1)
                r.font.color.rgb = RGBColor(0xA3, 0x15, 0x15)
            elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
                match = re.match(r'\[(.*?)\]\((.*?)\)', token)
                if match:
                    r = paragraph.add_run(match.group(1))
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(default_size)
                    r.font.color.rgb = RGBColor(0x00, 0x66, 0xCC)
                    r.underline = True
                else:
                    r = paragraph.add_run(token)
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(default_size)
            else:
                r = paragraph.add_run(token)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(default_size)

    i = 0
    while i < len(lines):
        line = lines[i]

        # Handle code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle tables
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()

        stripped = line.strip()

        # Empty lines
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CCCCCC"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # Headings (Times New Roman)
        if stripped.startswith('# '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[2:])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(20)
            r.bold = True
            r.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[3:])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(15)
            r.bold = True
            r.font.color.rgb = RGBColor(0x2B, 0x4C, 0x7E)
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(9)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[4:])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(0x3B, 0x62, 0x9B)
        elif stripped.startswith('#### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(7)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            r = h.add_run(stripped[5:])
            r.font.name = 'Times New Roman'
            r.font.size = Pt(11.5)
            r.bold = True
            r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
        # Bullet list items
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, stripped[2:])
        # Numbered list items
        elif re.match(r'^\d+\.\s+', stripped):
            match = re.match(r'^\d+\.\s+(.*)$', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, match.group(1))
        # Blockquotes
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="8" w:color="2B4C7E"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            add_formatted_text(p, stripped[2:])
        # Standard paragraph
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, stripped)

        i += 1

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == "__main__":
    doc_mappings = [
        ("docs/prd/PRD.md", "docs/prd/PRD.docx"),
        ("docs/trd/TRD.md", "docs/trd/TRD.docx"),
        ("docs/api-contracts/API_CONTRACT.md", "docs/api-contracts/API_CONTRACT.docx"),
        ("docs/database/NEO4J_SCHEMA.md", "docs/database/NEO4J_SCHEMA.docx")
    ]
    base_dir = "/home/yan2005dris-afk/Documentos/GitHub/RedSocial"
    for md_rel, docx_rel in doc_mappings:
        md_file = os.path.join(base_dir, md_rel)
        docx_file = os.path.join(base_dir, docx_rel)
        if os.path.exists(md_file):
            markdown_to_docx(md_file, docx_file)
